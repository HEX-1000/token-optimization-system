import hashlib
import re
import sys
import time
from pathlib import Path

from db import get_connection, init_db


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".json",
    ".yaml",
    ".yml",
    ".html",
    ".css",
    ".csv",
    ".docx",
    ".pdf",
    ".pptx",
    ".xlsx",
}


# ---------------------------------------------------------------------------
# File hashing
# ---------------------------------------------------------------------------

def file_hash(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)

            if not chunk:
                break

            digest.update(chunk)

    return digest.hexdigest()


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_plain_text(path: Path) -> str:
    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def extract_docx(path: Path):
    from docx import Document

    document = Document(path)
    items = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        if style_name.lower().startswith("heading"):
            items.append(
                (
                    "heading",
                    text,
                    style_name,
                )
            )
        else:
            items.append(
                (
                    "paragraph",
                    text,
                    style_name,
                )
            )

    return items


def extract_pdf(path: Path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []

    for index, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text() or ""

        pages.append(
            (
                index,
                text.strip(),
            )
        )

    return pages


def extract_pptx(path: Path):
    from pptx import Presentation

    presentation = Presentation(path)
    slides = []

    for index, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        parts = []
        title = ""

        if slide.shapes.title:
            title = (
                slide.shapes.title.text
                or ""
            ).strip()

        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue

            text = shape.text.strip()

            if not text:
                continue

            if text == title:
                continue

            parts.append(text)

        content_parts = []

        if title:
            content_parts.append(title)

        content_parts.extend(parts)

        content = "\n".join(
            content_parts
        ).strip()

        slides.append(
            (
                index,
                content,
                title,
            )
        )

    return slides


def extract_xlsx(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    sheets = []

    try:
        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                if any(
                    value.strip()
                    for value in values
                ):
                    rows.append(
                        " | ".join(values)
                    )

            content = "\n".join(
                rows
            ).strip()

            sheets.append(
                (
                    worksheet.title,
                    content,
                )
            )

    finally:
        workbook.close()

    return sheets


# ---------------------------------------------------------------------------
# Structural chunking
# ---------------------------------------------------------------------------

def chunk_markdown(text: str):
    """
    Split Markdown into structural chunks using ATX headings.

    Every heading starts a new chunk.

    Example:

        ## Section 1

        content...

        ## Section 2

        content...

    becomes:

        Section 1 -> heading + content
        Section 2 -> heading + content
    """

    lines = text.splitlines()

    chunks = []

    current_heading = None
    current_lines = []

    heading_pattern = re.compile(
        r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$"
    )

    def flush():
        nonlocal current_lines

        content = "\n".join(
            current_lines
        ).strip()

        if content:
            section = (
                current_heading
                if current_heading
                else "document: preamble"
            )

            chunks.append(
                (
                    section,
                    content,
                )
            )

        current_lines = []

    for line in lines:
        match = heading_pattern.match(line)

        if match:
            flush()

            current_heading = (
                match.group(2).strip()
            )

            current_lines = [
                line
            ]

        else:
            current_lines.append(line)

    flush()

    return chunks


def chunk_text(text: str):
    content = text.strip()

    if not content:
        return []

    return [
        (
            "document",
            content,
        )
    ]


def chunk_docx(items):
    chunks = []

    current_heading = None
    current_lines = []

    def flush():
        nonlocal current_lines

        content = "\n".join(
            current_lines
        ).strip()

        if content:
            section = (
                current_heading
                if current_heading
                else "document: preamble"
            )

            chunks.append(
                (
                    section,
                    content,
                )
            )

        current_lines = []

    for item_type, text, style_name in items:
        if item_type == "heading":
            flush()

            current_heading = text
            current_lines = [text]

        else:
            current_lines.append(text)

    flush()

    return chunks


# ---------------------------------------------------------------------------
# File -> chunks
# ---------------------------------------------------------------------------

def extract_chunks(path: Path):
    extension = path.suffix.lower()

    if extension == ".docx":
        return chunk_docx(
            extract_docx(path)
        )

    if extension == ".pdf":
        return [
            (
                f"page {page_number}",
                content,
            )
            for page_number, content
            in extract_pdf(path)
            if content
        ]

    if extension == ".pptx":
        return [
            (
                (
                    f"slide {slide_number}: {title}"
                    if title
                    else f"slide {slide_number}"
                ),
                content,
            )
            for slide_number, content, title
            in extract_pptx(path)
            if content
        ]

    if extension == ".xlsx":
        return [
            (
                f"sheet: {sheet_name}",
                content,
            )
            for sheet_name, content
            in extract_xlsx(path)
            if content
        ]

    text = extract_plain_text(path)

    if extension == ".md":
        return chunk_markdown(text)

    return chunk_text(text)


# ---------------------------------------------------------------------------
# Database helpers
# ---------------------------------------------------------------------------

def delete_file_chunks(conn, path: str):
    conn.execute(
        """
        DELETE FROM chunks
        WHERE path = ?
        """,
        (path,),
    )


def delete_root_index(conn, root: Path):
    root_string = str(root)

    rows = conn.execute(
        """
        SELECT path
        FROM files
        WHERE path = ?
           OR path LIKE ?
        """,
        (
            root_string,
            root_string + "\\%",
        ),
    ).fetchall()

    for row in rows:
        path = row["path"]

        delete_file_chunks(
            conn,
            path,
        )

        conn.execute(
            """
            DELETE FROM files
            WHERE path = ?
            """,
            (path,),
        )


def register_allowed_root(conn, root: Path):
    conn.execute(
        """
        INSERT OR IGNORE INTO allowed_roots (
            path,
            registered_at
        )
        VALUES (?, ?)
        """,
        (
            str(root),
            time.time(),
        ),
    )


def upsert_file(
    conn,
    path: Path,
    digest: str,
    chunks,
):
    stat = path.stat()

    existing = conn.execute(
        """
        SELECT *
        FROM files
        WHERE path = ?
        """,
        (str(path),),
    ).fetchone()

    now = time.time()

    if existing is None:
        conn.execute(
            """
            INSERT INTO files (
                path,
                name,
                extension,
                size,
                modified,
                content_hash,
                indexed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                str(path),
                path.name,
                path.suffix.lower(),
                stat.st_size,
                stat.st_mtime,
                digest,
                now,
            ),
        )

        return "new"

    if existing["content_hash"] == digest:
        if (
            existing["size"] != stat.st_size
            or existing["modified"] != stat.st_mtime
        ):
            conn.execute(
                """
                UPDATE files
                SET
                    name = ?,
                    extension = ?,
                    size = ?,
                    modified = ?,
                    indexed_at = ?
                WHERE path = ?
                """,
                (
                    path.name,
                    path.suffix.lower(),
                    stat.st_size,
                    stat.st_mtime,
                    now,
                    str(path),
                ),
            )

        return "skipped"

    conn.execute(
        """
        UPDATE files
        SET
            name = ?,
            extension = ?,
            size = ?,
            modified = ?,
            content_hash = ?,
            indexed_at = ?
        WHERE path = ?
        """,
        (
            path.name,
            path.suffix.lower(),
            stat.st_size,
            stat.st_mtime,
            digest,
            now,
            str(path),
        ),
    )

    return "updated"


def insert_chunks(
    conn,
    path: Path,
    chunks,
):
    for section, content in chunks:
        conn.execute(
            """
            INSERT INTO chunks (
                path,
                section,
                content
            )
            VALUES (?, ?, ?)
            """,
            (
                str(path),
                section,
                content,
            ),
        )


# ---------------------------------------------------------------------------
# Main indexing
# ---------------------------------------------------------------------------

def index_directory(
    root: Path,
    rebuild: bool = False,
):
    root = root.resolve()

    if not root.exists():
        raise FileNotFoundError(
            f"Directory does not exist: {root}"
        )

    if not root.is_dir():
        raise NotADirectoryError(
            f"Not a directory: {root}"
        )

    print(f"Allowed root: {root}")
    print(
        f"Mode: {'REBUILD' if rebuild else 'INCREMENTAL'}"
    )

    files = [
        path
        for path in root.rglob("*")
        if (
            path.is_file()
            and path.suffix.lower()
            in SUPPORTED_EXTENSIONS
        )
    ]

    files.sort()

    print(f"Found files: {len(files)}")
    print()

    conn = get_connection()

    new_count = 0
    updated_count = 0
    skipped_count = 0
    failed_count = 0
    deleted_count = 0

    try:
        register_allowed_root(
            conn,
            root,
        )

        if rebuild:
            delete_root_index(
                conn,
                root,
            )

        current_paths = {
            str(path.resolve())
            for path in files
        }

        indexed_rows = conn.execute(
            """
            SELECT path
            FROM files
            WHERE path = ?
               OR path LIKE ?
            """,
            (
                str(root),
                str(root) + "\\%",
            ),
        ).fetchall()

        indexed_paths = {
            row["path"]
            for row in indexed_rows
        }

        deleted_paths = (
            indexed_paths
            - current_paths
        )

        for deleted_path in sorted(
            deleted_paths
        ):
            print(
                f"Deleting missing: {deleted_path}"
            )

            delete_file_chunks(
                conn,
                deleted_path,
            )

            conn.execute(
                """
                DELETE FROM files
                WHERE path = ?
                """,
                (deleted_path,),
            )

            deleted_count += 1

        for path in files:
            try:
                digest = file_hash(path)

                existing = conn.execute(
                    """
                    SELECT content_hash
                    FROM files
                    WHERE path = ?
                    """,
                    (str(path),),
                ).fetchone()

                if (
                    not rebuild
                    and existing is not None
                    and existing["content_hash"] == digest
                ):
                    result = upsert_file(
                        conn,
                        path,
                        digest,
                        [],
                    )

                    if result == "skipped":
                        skipped_count += 1

                    continue

                print(
                    f"Indexing: {path}"
                )

                chunks = extract_chunks(path)

                result = upsert_file(
                    conn,
                    path,
                    digest,
                    chunks,
                )

                delete_file_chunks(
                    conn,
                    str(path),
                )

                insert_chunks(
                    conn,
                    path,
                    chunks,
                )

                if result == "new":
                    new_count += 1

                elif result == "updated":
                    updated_count += 1

                elif result == "skipped":
                    skipped_count += 1

            except Exception as exc:
                failed_count += 1

                print(
                    f"FAILED: {path}"
                )
                print(
                    f"  {type(exc).__name__}: {exc}"
                )

        conn.commit()

    finally:
        conn.close()

    print()
    print(f"New files: {new_count}")
    print(f"Updated files: {updated_count}")
    print(f"Skipped unchanged: {skipped_count}")
    print(f"Deleted files: {deleted_count}")
    print(f"Failed files: {failed_count}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    args = sys.argv[1:]

    if not args:
        print(
            "Usage: python indexer.py <directory> [--rebuild]"
        )
        raise SystemExit(1)

    root = Path(
        args[0]
    ).expanduser().resolve()

    rebuild = "--rebuild" in args[1:]

    unknown_args = [
        arg
        for arg in args[1:]
        if arg != "--rebuild"
    ]

    if unknown_args:
        print(
            "Unknown argument(s): "
            + " ".join(unknown_args)
        )
        print(
            "Usage: python indexer.py <directory> [--rebuild]"
        )
        raise SystemExit(1)

    init_db()

    index_directory(
        root,
        rebuild=rebuild,
    )


if __name__ == "__main__":
    main()


